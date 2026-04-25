from io import BytesIO

from openpyxl.reader.excel import load_workbook

from deps_output_exporting.infrastructure.services.document_layout.excel_builder import (
    ExcelBuilder,
    GeneralInfoFiller,
    KeyValuePairsFiller,
    ParagraphsFiller,
    PrompterKeyValuesFiller,
    TablesFiller,
)

FIRST_ELEMENT = 0
SECOND_ELEMENT = 1


def test_general_info_sheet_creation(excel_builder, test_workbook, general_info):
    excel_builder._create_general_info_sheet(test_workbook)

    assert (ws := test_workbook[ExcelBuilder.GENERAL_INFO_SHEET_NAME])
    assert ws["A1"].value == GeneralInfoFiller.GENERIC_INFO_VERTICAL_HEADER[0]
    assert ws["A2"].value == GeneralInfoFiller.GENERIC_INFO_VERTICAL_HEADER[1]
    assert ws["A3"].value == GeneralInfoFiller.GENERIC_INFO_VERTICAL_HEADER[2]
    assert ws["A4"].value == GeneralInfoFiller.GENERIC_INFO_VERTICAL_HEADER[3]
    assert ws["A5"].value == GeneralInfoFiller.GENERIC_INFO_VERTICAL_HEADER[4]
    assert ws["B1"].value == general_info.title
    assert ws["B2"].value == general_info.type
    assert ws["B3"].value == general_info.uploaded_date.strftime("%d-%m-%Y, %H:%M")
    assert ws["B4"].value == general_info.pages_number
    assert ws["B5"].value == general_info.engine


def test_paragraphs_sheet_creation(excel_builder, test_workbook, document_layout, page1, page2):
    excel_builder._create_paragraphs_sheet(test_workbook)

    assert (ws := test_workbook[ExcelBuilder.PARAGRAPH_SHEET_NAME])

    ws_rows = list(ws.iter_rows(values_only=True))

    assert ws_rows[0] == (ParagraphsFiller.PAGE_HEADER, page1.page_number)
    assert ws_rows[1] == (ParagraphsFiller.PARAGRAPHS_HEADER, None)
    assert ws_rows[2] == (page1.paragraphs[FIRST_ELEMENT].content, None)
    assert ws_rows[3] == (None, None)
    assert ws_rows[4] == (ParagraphsFiller.PAGE_HEADER, page2.page_number)
    assert ws_rows[5] == (ParagraphsFiller.PARAGRAPHS_HEADER, None)
    assert ws_rows[6] == (page2.paragraphs[FIRST_ELEMENT].content, None)


def test_kvp_sheet_creation(
    excel_builder,
    test_workbook,
    document_layout,
    page1,
    page2,
    kvp1,
    kvp2,
    kvp3,
    kvp4,
):
    excel_builder._create_key_value_pairs_sheet(test_workbook)

    assert (ws := test_workbook[ExcelBuilder.KV_DATA_SHEET_NAME])

    ws_rows = list(ws.iter_rows(values_only=True))

    assert ws_rows[0] == (KeyValuePairsFiller.PAGE_HEADER, page1.page_number, None)
    assert ws_rows[1] == tuple(KeyValuePairsFiller.KEY_VALUE_PAIR_HEADER)
    assert ws_rows[2] == (kvp1.key_content, kvp1.value_content, kvp1.confidence)
    assert ws_rows[3] == (kvp2.key_content, kvp2.value_content, kvp2.confidence)
    assert ws_rows[4] == (None, None, None)
    assert ws_rows[5] == (KeyValuePairsFiller.PAGE_HEADER, page2.page_number, None)
    assert ws_rows[6] == tuple(KeyValuePairsFiller.KEY_VALUE_PAIR_HEADER)
    assert ws_rows[7] == (kvp3.key_content, kvp3.value_content, kvp3.confidence)
    assert ws_rows[8] == (kvp4.key_content, kvp4.value_content, kvp4.confidence)


def test_table_view_sheet_creation(excel_builder, test_workbook, document_layout, page1, page2, table1, table2):
    excel_builder._create_table_views_sheet(test_workbook)

    assert (ws := test_workbook[ExcelBuilder.TABLE_VIEW_SHEET_NAME])

    ws_rows = list(ws.iter_rows(values_only=True))

    # fmt: off
    assert ws_rows[0] == (TablesFiller.PAGE_HEADER, page1.page_number, None, None)
    assert ws_rows[1] == (TablesFiller.TABLE_HEADER, None, None, None)
    assert ws_rows[2] == (
        TablesFiller.COLUMN_COUNT_HEADER, table1.column_count,
        TablesFiller.ROW_COUNT_HEADER, table1.row_count,
    )
    assert ws_rows[3] == (table1.cells[FIRST_ELEMENT].content, table1.cells[SECOND_ELEMENT].content, None, None)
    assert ws_rows[4] == (None, None, None, None)
    assert ws_rows[5] == (None, None, None, None)
    assert ws_rows[6] == (TablesFiller.PAGE_HEADER, page2.page_number, None, None)
    assert ws_rows[7] == (TablesFiller.TABLE_HEADER, None, None, None)
    assert ws_rows[8] == (
        TablesFiller.COLUMN_COUNT_HEADER, table2.column_count,
        TablesFiller.ROW_COUNT_HEADER, table2.row_count,
    )
    assert ws_rows[9] == (table2.cells[FIRST_ELEMENT].content, table2.cells[SECOND_ELEMENT].content, None, None)
    # fmt: on


def test_table_detail_sheet_creation(
    excel_builder,
    test_workbook,
    document_layout,
    page1,
    page2,
    table1,
    table2,
    cell1,
    cell2,
    cell3,
    cell4,
):
    excel_builder._create_table_details_sheet(test_workbook)

    assert (ws := test_workbook[ExcelBuilder.TABLE_DETAIL_SHEET_NAME])

    ws_rows = list(ws.iter_rows(values_only=True))

    # fmt: off
    assert ws_rows[0] == (TablesFiller.PAGE_HEADER, page1.page_number, None, None, None, None, None, None)
    assert ws_rows[1] == (TablesFiller.TABLE_HEADER, None, None, None, None, None, None, None)
    assert ws_rows[2] == (
        TablesFiller.COLUMN_COUNT_HEADER, table1.column_count,
        TablesFiller.ROW_COUNT_HEADER, table1.row_count,
        None, None, None, None,
    )
    assert ws_rows[3] == (
        TablesFiller.CONTENT_HEADER, cell1.content,
        TablesFiller.KIND_HEADER, cell1.kind,
        None, None, None, None,
    )
    assert ws_rows[4] == (
        TablesFiller.COLUMN_INDEX_HEADER, cell1.column_index,
        TablesFiller.ROW_INDEX_HEADER, cell1.row_index,
        TablesFiller.COLUMN_SPAN_HEADER, cell1.column_span,
        TablesFiller.ROW_SPAN_HEADER, cell1.row_span,
    )
    assert ws_rows[5] == (
        TablesFiller.CONTENT_HEADER, cell2.content,
        TablesFiller.KIND_HEADER, cell2.kind,
        None, None, None, None,
    )
    assert ws_rows[6] == (
        TablesFiller.COLUMN_INDEX_HEADER, cell2.column_index,
        TablesFiller.ROW_INDEX_HEADER, cell2.row_index,
        TablesFiller.COLUMN_SPAN_HEADER, cell2.column_span,
        TablesFiller.ROW_SPAN_HEADER, cell2.row_span,
    )

    assert ws_rows[7] == (None, None, None, None, None, None, None, None)
    assert ws_rows[8] == (None, None, None, None, None, None, None, None)

    assert ws_rows[9] == (TablesFiller.PAGE_HEADER, page2.page_number, None, None, None, None, None, None)
    assert ws_rows[10] == (TablesFiller.TABLE_HEADER, None, None, None, None, None, None, None)
    assert ws_rows[11] == (
        TablesFiller.COLUMN_COUNT_HEADER, table2.column_count,
        TablesFiller.ROW_COUNT_HEADER, table2.row_count,
        None, None, None, None,
    )
    assert ws_rows[12] == (
        TablesFiller.CONTENT_HEADER, cell3.content,
        TablesFiller.KIND_HEADER, cell3.kind,
        None, None, None, None,
    )
    assert ws_rows[13] == (
        TablesFiller.COLUMN_INDEX_HEADER, cell3.column_index,
        TablesFiller.ROW_INDEX_HEADER, cell3.row_index,
        TablesFiller.COLUMN_SPAN_HEADER, cell3.column_span,
        TablesFiller.ROW_SPAN_HEADER, cell3.row_span,
    )
    assert ws_rows[14] == (
        TablesFiller.CONTENT_HEADER, cell4.content,
        TablesFiller.KIND_HEADER, cell4.kind,
        None, None, None, None,
    )
    assert ws_rows[15] == (
        TablesFiller.COLUMN_INDEX_HEADER, cell4.column_index,
        TablesFiller.ROW_INDEX_HEADER, cell4.row_index,
        TablesFiller.COLUMN_SPAN_HEADER, cell4.column_span,
        TablesFiller.ROW_SPAN_HEADER, cell4.row_span,
    )
    # fmt: on


def test_excel_data_building__all_sheets_present(general_info, document_layout):
    data = ExcelBuilder.with_info(general_info=general_info, document_layout=document_layout).build()

    workbook = load_workbook(BytesIO(data))
    assert workbook.sheetnames == [
        ExcelBuilder.GENERAL_INFO_SHEET_NAME,
        ExcelBuilder.PARAGRAPH_SHEET_NAME,
        ExcelBuilder.KV_DATA_SHEET_NAME,
        ExcelBuilder.TABLE_VIEW_SHEET_NAME,
        ExcelBuilder.TABLE_DETAIL_SHEET_NAME,
    ]


def test_key_values_sheet_creation(excel_builder_with_key_values, test_workbook, key_values):
    excel_builder_with_key_values._create_prompter_key_values_sheet(test_workbook)

    assert (ws := test_workbook[ExcelBuilder.KEY_VALUES_SHEET_NAME])

    ws_rows = list(ws.iter_rows(values_only=True))
    assert ws_rows[0] == (PrompterKeyValuesFiller.KEY_VALUES_VERTICAL_HEADER[0], key_values[0].key)
    assert ws_rows[1] == (PrompterKeyValuesFiller.KEY_VALUES_VERTICAL_HEADER[1], key_values[0].value)
    assert ws_rows[3] == (PrompterKeyValuesFiller.KEY_VALUES_VERTICAL_HEADER[0], key_values[1].key)
    assert ws_rows[4] == (PrompterKeyValuesFiller.KEY_VALUES_VERTICAL_HEADER[1], key_values[1].value)


def test_excel_data_building__kvs_output_enabled__all_sheets_present(general_info, document_layout, key_values):
    data = (
        ExcelBuilder.with_info(general_info=general_info, document_layout=document_layout)
        .with_key_values(key_values)
        .build()
    )

    workbook = load_workbook(BytesIO(data))
    assert workbook.sheetnames == [
        ExcelBuilder.GENERAL_INFO_SHEET_NAME,
        ExcelBuilder.PARAGRAPH_SHEET_NAME,
        ExcelBuilder.KV_DATA_SHEET_NAME,
        ExcelBuilder.TABLE_VIEW_SHEET_NAME,
        ExcelBuilder.TABLE_DETAIL_SHEET_NAME,
        ExcelBuilder.KEY_VALUES_SHEET_NAME,
    ]
