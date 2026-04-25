import json
from io import BytesIO
from unittest.mock import ANY

import pytest
from openpyxl import load_workbook

from deps_output_exporting.domain.model import Output, OutputState
from deps_output_exporting.infrastructure.services import OutputBuildingService
from deps_output_exporting.infrastructure.services.document_layout.excel_builder import (
    ExcelBuilder,
)
from tests.data import (
    json_dl_output_all_features,
    json_dl_output_one_feature,
    json_dl_output_with_key_values,
)

FILEPATH = "output/"


def test_build_dl_json_output__all_features(
    output_building_service,
    test_tenant_id,
    profile__dl_json_all_features,
    document_type__dl_json_profiles,
    document_id,
    get_document_layout_all_features_success_request_mock,
    file_name_json,
    upload_json_file_mock,
    mocker,
):
    mocker.patch.object(OutputBuildingService, "_generate_file_name", return_value=file_name_json)

    output = output_building_service.build(
        document_id,
        test_tenant_id,
        profile__dl_json_all_features,
        document_type__dl_json_profiles.id(),
    )

    upload_json_file_mock.assert_called_with(
        FILEPATH, file_name_json, json.dumps(json_dl_output_all_features).encode("utf8")
    )

    assert isinstance(output, Output)
    assert output.tenant_id() == test_tenant_id
    assert output.profile_info.id == profile__dl_json_all_features.id
    assert output.profile_info.version == profile__dl_json_all_features.version
    assert output.document_id == document_id
    assert output.state == OutputState.READY
    assert output.file_path == f"{FILEPATH}{file_name_json}"
    assert output.creation_date


def test_build_dl_json_output__one_feature(
    output_building_service,
    test_tenant_id,
    profile__dl_json_one_feature,
    document_type__dl_json_profiles,
    document_id,
    get_document_layout_kvp_feature_success_request_mock,
    file_name_json,
    upload_json_file_mock,
    mocker,
):
    mocker.patch.object(OutputBuildingService, "_generate_file_name", return_value=file_name_json)

    output = output_building_service.build(
        document_id,
        test_tenant_id,
        profile__dl_json_one_feature,
        document_type__dl_json_profiles.id(),
    )

    upload_json_file_mock.assert_called_with(
        FILEPATH, file_name_json, json.dumps(json_dl_output_one_feature).encode("utf8")
    )

    assert isinstance(output, Output)
    assert output.tenant_id() == test_tenant_id
    assert output.profile_info.id == profile__dl_json_one_feature.id
    assert output.profile_info.version == profile__dl_json_one_feature.version
    assert output.document_id == document_id
    assert output.state == OutputState.READY
    assert output.file_path == f"{FILEPATH}{file_name_json}"
    assert output.creation_date


@pytest.mark.usefixtures("enable_key_values_output")
def test_build_dl_json_output__kvs_output_enabled(
    output_building_service,
    test_tenant_id,
    profile__dl_json_all_features,
    document_type__dl_json_profiles,
    document_id,
    get_document_layout_all_features_success_request_mock,
    get_prompter_key_values_success_request_mock,
    file_name_json,
    upload_json_file_mock,
    mocker,
):
    mocker.patch.object(OutputBuildingService, "_generate_file_name", return_value=file_name_json)

    output_building_service.build(
        document_id,
        test_tenant_id,
        profile__dl_json_all_features,
        document_type__dl_json_profiles.id(),
    )

    upload_json_file_mock.assert_called_with(
        FILEPATH, file_name_json, json.dumps(json_dl_output_with_key_values).encode("utf8")
    )


def test_build_dl_excel_output__ok(
    output_building_service,
    test_tenant_id,
    profile__dl_excel_all_features,
    document_type__dl_excel_profiles,
    document_id,
    get_document_layout_all_features_success_request_mock,
    get_document_type_success_request_mock,
    get_document_success_request_mock,
    file_name_json,
    upload_json_file_mock,
    mocker,
):
    mocker.patch.object(OutputBuildingService, "_generate_file_name", return_value=file_name_json)

    output = output_building_service.build(
        document_id,
        test_tenant_id,
        profile__dl_excel_all_features,
        document_type__dl_excel_profiles.id(),
    )

    upload_json_file_mock.assert_called_with(FILEPATH, file_name_json, ANY)

    assert isinstance(output, Output)
    assert output.tenant_id() == test_tenant_id
    assert output.profile_info.id == profile__dl_excel_all_features.id
    assert output.profile_info.version == profile__dl_excel_all_features.version
    assert output.document_id == document_id
    assert output.state == OutputState.READY
    assert output.file_path == f"{FILEPATH}{file_name_json}"
    assert output.creation_date


def test_build_dl_excel_output__data_created(
    get_document_layout_all_features_success_request_mock,
    get_document_type_success_request_mock,
    get_document_success_request_mock,
    document_type__dl_excel_profiles,
    profile__dl_excel_all_features,
    output_building_service,
    upload_excel_file_mock,
    file_name_excel,
    document_layout,
    test_tenant_id,
    general_info,
    document_id,
    mocker,
):
    mocker.patch.object(OutputBuildingService, "_generate_file_name", return_value=file_name_excel)
    create_excel_data = mocker.patch.object(ExcelBuilder, "with_info")

    output_building_service.build(
        document_id,
        test_tenant_id,
        profile__dl_excel_all_features,
        document_type__dl_excel_profiles.id(),
    )

    create_excel_data.assert_called_with(general_info=general_info, document_layout=document_layout)


def test_build_dl_excel_output__engine_in_document_detail__passed_to_builder(
    get_document_layout_all_features_success_request_mock,
    get_document_type_success_request_mock,
    get_document_success_request_mock,
    document_type__dl_excel_profiles,
    profile__dl_excel_all_features,
    output_building_service,
    upload_excel_file_mock,
    file_name_excel,
    test_tenant_id,
    document_id,
    mocker,
):
    mocker.patch.object(OutputBuildingService, "_generate_file_name", return_value=file_name_excel)
    create_excel_data = mocker.patch.object(ExcelBuilder, "with_info")

    output_building_service.build(
        document_id,
        test_tenant_id,
        profile__dl_excel_all_features,
        document_type__dl_excel_profiles.id(),
    )

    assert create_excel_data.call_args.kwargs["general_info"].engine == "TESSERACT"


def test_build_dl_excel_output__engine_in_document_type__passed_to_builder(
    get_document_layout_all_features_success_request_mock,
    get_document_type_success_request_mock,
    get_document_success_request_mock__no_engine,
    document_type__dl_excel_profiles,
    profile__dl_excel_all_features,
    output_building_service,
    upload_excel_file_mock,
    file_name_excel,
    test_tenant_id,
    document_id,
    mocker,
):
    mocker.patch.object(OutputBuildingService, "_generate_file_name", return_value=file_name_excel)
    create_excel_data = mocker.patch.object(ExcelBuilder, "with_info")

    output_building_service.build(
        document_id,
        test_tenant_id,
        profile__dl_excel_all_features,
        document_type__dl_excel_profiles.id(),
    )

    assert create_excel_data.call_args.kwargs["general_info"].engine == "TESSERACT"


def test_build_dl_excel_output__default_engine__passed_to_builder(
    get_document_layout_all_features_success_request_mock,
    get_document_type_success_request_mock__no_engine,
    get_document_success_request_mock__no_engine,
    document_type__dl_excel_profiles,
    profile__dl_excel_all_features,
    output_building_service,
    upload_excel_file_mock,
    file_name_excel,
    test_tenant_id,
    document_id,
    mocker,
):
    mocker.patch.object(OutputBuildingService, "_generate_file_name", return_value=file_name_excel)
    create_excel_data = mocker.patch.object(ExcelBuilder, "with_info")

    output_building_service.build(
        document_id,
        test_tenant_id,
        profile__dl_excel_all_features,
        document_type__dl_excel_profiles.id(),
    )

    assert create_excel_data.call_args.kwargs["general_info"].engine == "TESSERACT"


@pytest.mark.usefixtures("enable_key_values_output")
def test_build_dl_excel_output__kvs_output_enabled(
    get_document_layout_all_features_success_request_mock,
    get_prompter_key_values_success_request_mock,
    get_document_type_success_request_mock,
    get_document_success_request_mock,
    document_type__dl_excel_profiles,
    profile__dl_excel_all_features,
    output_building_service,
    upload_excel_file_mock,
    file_name_excel,
    key_values,
    test_tenant_id,
    document_id,
    mocker,
):
    mocker.patch.object(OutputBuildingService, "_generate_file_name", return_value=file_name_excel)
    create_excel_kvs_data = mocker.patch.object(ExcelBuilder, "with_key_values")

    output_building_service.build(
        document_id,
        test_tenant_id,
        profile__dl_excel_all_features,
        document_type__dl_excel_profiles.id(),
    )

    create_excel_kvs_data.assert_called_with(key_values)


def test_build_dl_with_merged_tables__json_output__ok(
    output_building_service,
    test_tenant_id,
    profile__dl_json_all_features,
    document_type__dl_excel_profiles,
    document_id,
    get_document_layout_with_merged_tables_all_features_success_request_mock,
    get_document_type_success_request_mock,
    get_document_success_request_mock,
    file_name_json,
    upload_json_file_mock,
    mocker,
):
    mocker.patch.object(OutputBuildingService, "_generate_file_name", return_value=file_name_json)
    output_building_service.build(
        document_id,
        test_tenant_id,
        profile__dl_json_all_features,
        document_type__dl_excel_profiles.id(),
    )

    upload_json_file_mock.assert_called_with(FILEPATH, file_name_json, ANY)

    json_output_data = json.loads(upload_json_file_mock.call_args.args[2])
    expected_len_of_tables_per_page = (1, 0, 1, 1, 0, 1)
    for exp_length, page in zip(expected_len_of_tables_per_page, json_output_data["pages"]):
        assert exp_length == len(page["tables"])


def test_build_dl_with_merged_tables__excel_output__ok(
    output_building_service,
    test_tenant_id,
    profile__dl_excel_all_features,
    document_type__dl_excel_profiles,
    document_id,
    get_document_layout_with_merged_tables_all_features_success_request_mock,
    get_document_type_success_request_mock,
    get_document_success_request_mock,
    file_name_json,
    upload_json_file_mock,
    mocker,
):
    mocker.patch.object(OutputBuildingService, "_generate_file_name", return_value=file_name_json)

    output_building_service.build(
        document_id,
        test_tenant_id,
        profile__dl_excel_all_features,
        document_type__dl_excel_profiles.id(),
    )

    upload_json_file_mock.assert_called_with(FILEPATH, file_name_json, ANY)

    counts_rows = _get_counts_rows_from_excel(upload_json_file_mock.call_args.args[2])
    expected_table_counts = ((5, 5), (3, 7), (2, 10), (6, 6))

    assert len(counts_rows) == 4
    for exp_result, gotten_res in zip(expected_table_counts, counts_rows):
        assert exp_result == tuple(gotten_res)


def _get_counts_rows_from_excel(data: bytes) -> list[list[int]]:
    wb = load_workbook(filename=BytesIO(data))
    sheet = wb["Table View"]

    counts_rows = []
    for row in sheet.iter_rows():
        if (value := row[0].value) is not None and value.startswith("Column count"):
            counts_rows.append([cell.value for cell in row if isinstance(cell.value, int)])
    wb.close()
    return counts_rows
