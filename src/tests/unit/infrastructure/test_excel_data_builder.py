from io import BytesIO

from openpyxl import load_workbook

from deps_output_exporting.infrastructure.services import ExcelDataBuilder


class TestExcelFileBuilder:
    def test_general_info_sheet_creation(self, excel_data_builder, test_workbook):
        excel_data_builder._create_general_info_sheet(test_workbook)

        assert (ws := test_workbook["General Information"])
        assert ws["A1"].value == "Document Name"
        assert ws["A2"].value == "Type"
        assert ws["A3"].value == "Uploaded Date"
        assert ws["A4"].value == "Number of Fields"
        assert ws["A5"].value == "OCR Engine"
        assert ws["B1"].value == "Document Title"
        assert ws["B2"].value == "AllFieldTypesQA"
        assert ws["B3"].value == "11-12-2023, 11:52"
        assert ws["B4"].value == 18
        assert ws["B5"].value == "TESSERACT"

    def test_generic_data_sheet_creation(self, excel_data_builder, test_workbook):
        excel_data_builder._create_generic_data_sheet(test_workbook)

        assert (ws := test_workbook["Generic Data"])

        ws_rows = list(ws.iter_rows(values_only=True))

        assert ws_rows[0] == ("Name", "Page", "Type", "Value", "Confidence", "Validation")
        assert ws_rows[1] == ("Checkbox 1", 1, "Checkbox", True, "0.97", "Failed")
        assert ws_rows[2] == ("Enum 1", 1, "Enum", "3b65e7c495004646bd567e333946154c", "0.86", "Passed")
        assert ws_rows[3] == ("String 1", 2, "String", "Bunting", "0.92", "Passed")

    def test_kv_data_sheet_creation(self, excel_data_builder, test_workbook):
        excel_data_builder._create_kv_data_sheet(test_workbook)

        assert (ws := test_workbook["Key Value Pairs"])

        ws_rows = list(ws.iter_rows(values_only=True))

        assert ws_rows[0] == (
            "Name",
            "Page",
            "Key",
            "Value",
            "Key Confidence",
            "Value Confidence",
            "Key Validation",
            "Value Validation",
        )
        assert ws_rows[1] == (
            "Key-Value Pair 1",
            1,
            "985a37dfb6b34d8082e6157c7e1b86fd",
            "0c197676253f44a89f882eb7296b1861",
            "0.99",
            "0.86",
            "Failed",
            "Failed",
        )

    def test_table_view_sheet_creation(self, excel_data_builder, test_workbook):
        excel_data_builder._create_table_view_sheet(test_workbook)

        assert (ws := test_workbook["Table View"])

        ws_rows = list(ws.iter_rows(values_only=True))

        assert ws_rows[0] == ("Field name", "Table 1", "Page", 2)
        assert ws_rows[1] == ("Col1", "Col2", None, None)
        assert ws_rows[2] == ("dfdbc8f0bf624bb7b4cccfc65d01a060", None, None, None)

    def test_table_detail_sheet_creation(self, excel_data_builder, test_workbook):
        excel_data_builder._create_table_detail_sheet(test_workbook)

        assert (ws := test_workbook["Table Detail"])

        ws_rows = list(ws.iter_rows(values_only=True))

        assert ws_rows[0] == ("Field name", "Table 1", "Page", 2, None)
        assert ws_rows[1] == ("Value", "Confidence", "Validation", "Column index", "Row index")
        assert ws_rows[2] == ("dfdbc8f0bf624bb7b4cccfc65d01a060", "0.80", "Passed", 0, 0)

    def test_list_data_sheet_creation(self, excel_data_builder, test_workbook):
        excel_data_builder._create_list_data_sheet(test_workbook)

        assert (ws := test_workbook["List Data"])

        ws_rows_values = list(ws.iter_rows(values_only=True))
        ws_rows_cells = list(ws.iter_rows())

        self._assert_generic_list_data(ws_rows_values)
        self._assert_kv_list_data(ws_rows_values)
        self._assert_table_list_data(ws_rows_values, ws_rows_cells, ws)

    def _assert_generic_list_data(self, ws_rows_values):
        assert ws_rows_values[0] == ("Field name", "Checkbox list 1", "Page", 1, "Type", "Checkbox", None)
        assert ws_rows_values[1] == ("Index", "Value", "Confidence", "Validation", None, None, None)
        assert ws_rows_values[2] == (1, True, "0.98", "Passed", None, None, None)
        assert ws_rows_values[3] == (2, False, "0.81", "Passed", None, None, None)
        assert ws_rows_values[4] == (3, None, "0.86", "Passed", None, None, None)

        assert ws_rows_values[6] == ("Field name", "Enum list 1", "Page", 1, "Type", "Enum", None)
        assert ws_rows_values[7] == ("Index", "Value", "Confidence", "Validation", None, None, None)
        assert ws_rows_values[8] == (1, "48262c4c407a45f098195bf293b3de28", "0.87", "Passed", None, None, None)
        assert ws_rows_values[9] == (2, "f", "0.92", "Passed", None, None, None)

        assert ws_rows_values[11] == ("Field name", "String list 1", "Page", 2, "Type", "String", None)
        assert ws_rows_values[12] == ("Index", "Value", "Confidence", "Validation", None, None, None)
        assert ws_rows_values[13] == (1, "780bc476656c4597916b367287e0105f", "0.86", "Passed", None, None, None)
        assert ws_rows_values[14] == (2, "0cd034a6b87242b5a0f1d8f8e624ac03", "0.84", "Failed", None, None, None)
        assert ws_rows_values[15] == (3, "bea8da5b61e842a581e3a3e1b3f3fac8", "0.89", "Passed", None, None, None)

    def _assert_kv_list_data(self, ws_rows_values):
        assert ws_rows_values[29] == ("Field name", "Key-Value Pair list 1", "Page", 1, "Type", "Key-Value Pair", None)
        assert ws_rows_values[30] == (
            "Index",
            "Key",
            "Value",
            "Key Confidence",
            "Value Confidence",
            "Key Validation",
            "Value Validation",
        )
        assert ws_rows_values[31] == (
            1,
            "c7ebd9f27d834143b8d809aa42ea20ca",
            "9bca244a30ff4a02a995f0f870b66d30",
            "0.94",
            "0.92",
            "Passed",
            "Failed",
        )
        assert ws_rows_values[32] == (
            2,
            "564b27777f1b41949bad51bdbd573f51",
            "4bf46332fb454229b6428941242c48e8",
            "0.91",
            "0.93",
            "Failed",
            "Passed",
        )

    def _assert_table_list_data(self, ws_rows_values, ws_rows_cells, ws):
        assert ws_rows_values[38] == ("Field name", "Table list 1", "Page", 2, "Type", "Table", None)
        assert ws_rows_values[39] == ("Validation", "Failed", None, None, None, None, None)
        assert ws_rows_values[40] == ("Index", "Table view", "Table detail", None, None, None, None)
        assert ws_rows_cells[41][0].value == 1
        assert ws_rows_cells[41][1].hyperlink.target == f"#'List Data'!{ws.tables['Table_list_1_0_view'].ref}"
        assert ws_rows_cells[41][2].hyperlink.target == f"#'List Data'!{ws.tables['Table_list_1_0_detail'].ref}"
        assert ws_rows_cells[42][0].value == 2
        assert ws_rows_cells[42][1].hyperlink.target == f"#'List Data'!{ws.tables['Table_list_1_1_view'].ref}"
        assert ws_rows_cells[42][2].hyperlink.target == f"#'List Data'!{ws.tables['Table_list_1_1_detail'].ref}"

        assert ws_rows_values[44] == ("Col1", "Col2", None, None, None, None, None)
        assert ws_rows_values[45] == ("6ac94d17d3764f9583ed9325d40e3948", None, None, None, None, None, None)

        assert ws_rows_values[47] == ("Col1", "Col2", None, None, None, None, None)
        assert ws_rows_values[48] == ("16fa5884783647e6a4e1dc0f46b0c330", None, None, None, None, None, None)

        assert ws_rows_values[50] == ("Value", "Confidence", "Validation", "Column index", "Row index", None, None)
        assert ws_rows_values[51] == ("6ac94d17d3764f9583ed9325d40e3948", "0.86", "Passed", 0, 0, None, None)

        assert ws_rows_values[53] == ("Value", "Confidence", "Validation", "Column index", "Row index", None, None)
        assert ws_rows_values[54] == ("16fa5884783647e6a4e1dc0f46b0c330", "0.98", "Failed", 0, 0, None, None)

    def test_table_view_sheet_creation__correct_cells_position(
        self, general_info, fields_data__table_field, test_workbook
    ):
        excel_data_builder = ExcelDataBuilder(general_info=general_info, fields_data=fields_data__table_field)
        excel_data_builder._create_table_view_sheet(test_workbook)

        ws_rows = list(test_workbook["Table View"].iter_rows(values_only=True))

        assert ws_rows[0] == ("Field name", "Table 1", "Page", 1)
        assert ws_rows[1] == ("Col1", "Col2", "Col3", None)
        assert ws_rows[2] == ("val1", "val2", "val3", None)
        assert ws_rows[3] == ("val4", "val5", "val6", None)

    def test_excel_data_building__all_field_types__all_sheets_present(
        self, general_info, fields_data__all_fields_with_validation
    ):
        data = ExcelDataBuilder.with_info(
            general_info=general_info,
            fields_data=fields_data__all_fields_with_validation,
        ).build()
        workbook = load_workbook(BytesIO(data))
        assert workbook.sheetnames == [
            "General Information",
            "Generic Data",
            "Key Value Pairs",
            "Table View",
            "Table Detail",
            "List Data",
        ]

    def test_excel_data_building__some_field_types__only_filled_sheets_present(
        self, general_info, fields_data__some_fields_without_validation
    ):
        data = ExcelDataBuilder.with_info(
            general_info=general_info,
            fields_data=fields_data__some_fields_without_validation,
        ).build()
        workbook = load_workbook(BytesIO(data))
        assert workbook.sheetnames == [
            "General Information",
            "Generic Data",
            "Key Value Pairs",
            "Table View",
            "Table Detail",
        ]
