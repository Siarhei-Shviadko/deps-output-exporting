from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from deps_output_exporting.infrastructure.proxies import Table as DocumentLayoutTable

from .base_filler import BaseFiller

__all__ = ["TablesFiller"]


class TablesFiller(BaseFiller):
    TABLE_HEADER = "Table"
    COLUMN_COUNT_HEADER = "Column count"
    ROW_COUNT_HEADER = "Row count"
    CONTENT_HEADER = "Content"
    KIND_HEADER = "Kind"
    COLUMN_INDEX_HEADER = "Column index"
    ROW_INDEX_HEADER = "Row index"
    COLUMN_SPAN_HEADER = "Column span"
    ROW_SPAN_HEADER = "Row span"
    HEADER_INDICATION = "header"

    def fill_tables_view_sheet(
        self,
        worksheet: Worksheet,
        page_numbers_tables: dict[int, list[DocumentLayoutTable]],
    ) -> None:
        for page_number, tables in page_numbers_tables.items():
            self._add_page_header(worksheet=worksheet, page_number=page_number)
            for table in tables:
                self._add_table_header(worksheet=worksheet)
                self._add_column_and_row_counts(worksheet=worksheet, table=table)
                self._add_cell_views(worksheet=worksheet, table=table, start_row=worksheet.max_row + 1)
                self._add_empty_row(worksheet)

            self._add_empty_row(worksheet)

    def fill_tables_detail_sheet(
        self,
        worksheet: Worksheet,
        page_numbers_tables: dict[int, list[DocumentLayoutTable]],
    ) -> None:
        for page_number, tables in page_numbers_tables.items():
            self._add_page_header(worksheet=worksheet, page_number=page_number)
            for table in tables:
                self._add_table_header(worksheet=worksheet)
                self._add_column_and_row_counts(worksheet=worksheet, table=table)
                self._add_cell_details(worksheet, table)
                self._add_empty_row(worksheet)

            self._add_empty_row(worksheet)

    def _add_cell_views(self, worksheet: Worksheet, table: DocumentLayoutTable, start_row: int) -> None:
        for cell in table.cells:
            ws_cell = worksheet.cell(
                row=cell.row_index + start_row,
                column=cell.column_index + 1,
                value=cell.content,
            )

            if cell.kind and self.HEADER_INDICATION in cell.kind.lower():
                ws_cell.font = Font(bold=True)

            if (rowspan := cell.row_span) > 1 and (colspan := cell.column_span) > 1:
                worksheet.merge_cells(
                    start_row=ws_cell.row,
                    start_column=ws_cell.column,
                    end_row=ws_cell.row + (rowspan - 1),
                    end_column=ws_cell.column + (colspan - 1),
                )
            elif (rowspan := cell.row_span) > 1:
                worksheet.merge_cells(
                    start_row=ws_cell.row,
                    start_column=ws_cell.column,
                    end_row=ws_cell.row + (rowspan - 1),
                    end_column=ws_cell.column,
                )
            elif (colspan := cell.column_span) > 1:
                worksheet.merge_cells(
                    start_row=ws_cell.row,
                    start_column=ws_cell.column,
                    end_row=ws_cell.row,
                    end_column=ws_cell.column + (colspan - 1),
                )

    def _add_cell_details(self, worksheet: Worksheet, table: DocumentLayoutTable) -> None:
        for cell in table.cells:
            worksheet.append([self.CONTENT_HEADER, cell.content, self.KIND_HEADER, cell.kind])
            self._add_bold_style_to_odd_cells(worksheet, worksheet.max_row)
            worksheet.append(
                [
                    self.COLUMN_INDEX_HEADER,
                    cell.column_index,
                    self.ROW_INDEX_HEADER,
                    cell.row_index,
                    self.COLUMN_SPAN_HEADER,
                    cell.column_span,
                    self.ROW_SPAN_HEADER,
                    cell.row_span,
                ],
            )
            self._add_bold_style_to_odd_cells(worksheet, worksheet.max_row)

    def _add_table_header(self, worksheet: Worksheet) -> None:
        worksheet.append([self.TABLE_HEADER])
        self._add_bold_style_to_row(worksheet, worksheet.max_row)

    def _add_column_and_row_counts(self, worksheet: Worksheet, table: DocumentLayoutTable) -> None:
        worksheet.append([self.COLUMN_COUNT_HEADER, table.column_count, self.ROW_COUNT_HEADER, table.row_count])
        self._add_bold_style_to_odd_cells(worksheet, worksheet.max_row)
