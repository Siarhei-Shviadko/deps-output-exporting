from io import BytesIO
from typing import Optional

from openpyxl.workbook import Workbook

from deps_output_exporting.infrastructure.proxies import DocumentLayout, KeyValues

from ..general_info import GeneralInfo
from .general_info_filler import GeneralInfoFiller
from .key_value_pairs_filler import KeyValuePairsFiller
from .paragraphs_filler import ParagraphsFiller
from .prompter_key_values_filler import PrompterKeyValuesFiller
from .tables_filler import TablesFiller

__all__ = ["ExcelBuilder"]


class ExcelBuilder:
    MIN_COLUMN_WIDTH = 6
    MAX_COLUMN_WIDTH = 150
    ADDITIONAL_WIDTH = 2

    GENERAL_INFO_SHEET_NAME = "General Information"
    PARAGRAPH_SHEET_NAME = "Paragraphs"
    KV_DATA_SHEET_NAME = "Key Value Pairs"
    TABLE_VIEW_SHEET_NAME = "Table View"
    TABLE_DETAIL_SHEET_NAME = "Table Detail"

    KEY_VALUES_SHEET_NAME = "Key Values"

    def __init__(self, general_info: GeneralInfo, document_layout: DocumentLayout) -> None:
        self._general_info_filler: GeneralInfoFiller = GeneralInfoFiller()
        self._paragraphs_filler: ParagraphsFiller = ParagraphsFiller()
        self._key_value_pairs_filler: KeyValuePairsFiller = KeyValuePairsFiller()
        self._tables_filler: TablesFiller = TablesFiller()
        self._prompter_key_values_filler: PrompterKeyValuesFiller = PrompterKeyValuesFiller()

        self._general_info = general_info
        self._document_layout = document_layout

        self._key_values: Optional[KeyValues] = None

    @classmethod
    def with_info(cls, general_info: GeneralInfo, document_layout: DocumentLayout) -> "ExcelBuilder":
        return cls(general_info, document_layout)

    def with_key_values(self, key_values: list[KeyValues]) -> "ExcelBuilder":
        self._key_values = key_values
        return self

    def build(self) -> bytes:
        workbook = Workbook()
        active = workbook.active
        workbook.remove(active)

        self._create_general_info_sheet(workbook)
        if self._document_layout.paragraphs:
            self._create_paragraphs_sheet(workbook)
        if self._document_layout.key_value_pairs:
            self._create_key_value_pairs_sheet(workbook)
        if self._document_layout.tables:
            self._create_table_views_sheet(workbook)
            self._create_table_details_sheet(workbook)

        if self._key_values is not None:
            self._create_prompter_key_values_sheet(workbook)

        self._set_column_width(workbook)

        with BytesIO() as buffer:
            workbook.save(buffer)
            return buffer.getvalue()

    def _create_general_info_sheet(self, workbook: Workbook) -> None:
        worksheet = workbook.create_sheet(self.GENERAL_INFO_SHEET_NAME)
        self._general_info_filler.add_general_info(worksheet=worksheet, general_info=self._general_info)

    def _create_paragraphs_sheet(self, workbook: Workbook) -> None:
        worksheet = workbook.create_sheet(self.PARAGRAPH_SHEET_NAME)
        self._paragraphs_filler.add_paragraphs(
            worksheet=worksheet,
            page_numbers_paragraphs=self._document_layout.paragraphs,
        )

    def _create_key_value_pairs_sheet(self, workbook: Workbook) -> None:
        worksheet = workbook.create_sheet(self.KV_DATA_SHEET_NAME)
        self._key_value_pairs_filler.add_key_value_pairs(
            worksheet=worksheet,
            page_numbers_key_value_pairs=self._document_layout.key_value_pairs,
        )

    def _create_table_views_sheet(self, workbook: Workbook) -> None:
        worksheet = workbook.create_sheet(self.TABLE_VIEW_SHEET_NAME)
        self._tables_filler.fill_tables_view_sheet(
            worksheet=worksheet,
            page_numbers_tables=self._document_layout.tables,
        )

    def _create_table_details_sheet(self, workbook: Workbook) -> None:
        worksheet = workbook.create_sheet(self.TABLE_DETAIL_SHEET_NAME)
        self._tables_filler.fill_tables_detail_sheet(
            worksheet=worksheet,
            page_numbers_tables=self._document_layout.tables,
        )

    def _create_prompter_key_values_sheet(self, workbook: Workbook) -> None:
        worksheet = workbook.create_sheet(self.KEY_VALUES_SHEET_NAME)
        self._prompter_key_values_filler.fill_key_values_sheet(worksheet, key_values=self._key_values)

    def _set_column_width(self, workbook: Workbook) -> None:
        for worksheet in workbook.worksheets:
            for column_cells in worksheet.columns:
                content_length_iterator = (len(str(value)) for cell in column_cells if (value := cell.value))
                max_content_length = max(content_length_iterator) + self.ADDITIONAL_WIDTH
                length_to_set = min(max_content_length, self.MAX_COLUMN_WIDTH)
                length_to_set = max(length_to_set, self.MIN_COLUMN_WIDTH)
                column_letter = column_cells[0].coordinate[0]
                worksheet.column_dimensions[column_letter].width = length_to_set
