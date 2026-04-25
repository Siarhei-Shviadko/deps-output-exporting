from openpyxl.worksheet.worksheet import Worksheet

from deps_output_exporting.infrastructure.proxies import Paragraph

from .base_filler import BaseFiller

__all__ = ["ParagraphsFiller"]


class ParagraphsFiller(BaseFiller):
    PARAGRAPHS_HEADER = "Paragraphs"

    def add_paragraphs(self, worksheet: Worksheet, page_numbers_paragraphs: dict[int, list[Paragraph]]) -> None:
        for page_number, paragraphs in page_numbers_paragraphs.items():
            self._add_page_header(worksheet=worksheet, page_number=page_number)
            self._add_paragraph_header(worksheet)

            for paragraph in paragraphs:
                worksheet.append([paragraph.content])

            self._add_empty_row(worksheet)

    def _add_paragraph_header(self, worksheet: Worksheet) -> None:
        worksheet.append([self.PARAGRAPHS_HEADER])
        self._add_bold_style_to_row(worksheet, worksheet.max_row)
