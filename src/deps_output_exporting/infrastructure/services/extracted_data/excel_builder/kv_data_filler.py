from openpyxl.worksheet.worksheet import Worksheet

from deps_output_exporting.infrastructure.proxies import ExtractedField

from .base_data_filler import BaseDataFiller
from .confidence_calculator import get_confidence_table_representation

__all__ = ["KeyValueDataFiller"]


class KeyValueDataFiller(BaseDataFiller):
    KV_DATA_TABLE_HEADER = [
        "Name",
        "Page",
        "Key",
        "Value",
        "Key Confidence",
        "Value Confidence",
        "Key Validation",
        "Value Validation",
    ]
    LIST_KV_DATA_TABLE_HEADER = [
        "Index",
        "Key",
        "Value",
        "Key Confidence",
        "Value Confidence",
        "Key Validation",
        "Value Validation",
    ]

    def add_kv_data(self, worksheet: Worksheet, kv_fields: list[ExtractedField]) -> None:
        worksheet.append(self.KV_DATA_TABLE_HEADER)
        self._add_bold_style_to_row(worksheet, worksheet.max_row)
        for field in kv_fields:
            worksheet.append(
                [
                    field.name,
                    field.page,
                    field.data.key.value,
                    field.data.value.value,
                    get_confidence_table_representation(field.data.key.confidence),
                    get_confidence_table_representation(field.data.value.confidence),
                    field.data.key.validation_result.value,
                    field.data.value.validation_result.value,
                ],
            )

    def add_kv_list_fields_data(self, worksheet: Worksheet, kv_list_fields: list[ExtractedField]) -> None:
        for kv_list_field in kv_list_fields:
            worksheet.append(
                self._create_list_common_header(
                    kv_list_field.name,
                    kv_list_field.page,
                    self.TYPES_REPRESENTATION[kv_list_field.base_type],
                ),
            )
            self._add_bold_style_to_odd_cells(worksheet, worksheet.max_row)
            worksheet.append(self.LIST_KV_DATA_TABLE_HEADER)
            self._add_bold_style_to_row(worksheet, worksheet.max_row)
            for index, item in enumerate(kv_list_field.data, 1):
                worksheet.append(
                    [
                        index,
                        item.key.value,
                        item.value.value,
                        get_confidence_table_representation(item.key.confidence),
                        get_confidence_table_representation(item.value.confidence),
                        item.key.validation_result.value,
                        item.value.validation_result.value,
                    ],
                )
            worksheet.append([])
