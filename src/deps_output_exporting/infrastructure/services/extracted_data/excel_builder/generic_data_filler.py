from openpyxl.worksheet.worksheet import Worksheet

from deps_output_exporting.infrastructure.proxies import ExtractedField

from .base_data_filler import BaseDataFiller
from .confidence_calculator import get_confidence_table_representation

__all__ = ["GenericDataFiller"]


class GenericDataFiller(BaseDataFiller):
    GENERIC_DATA_TABLE_HEADER = ["Name", "Page", "Type", "Value", "Confidence", "Validation"]
    LIST_GENERIC_DATA_TABLE_HEADER = ["Index", "Value", "Confidence", "Validation"]

    def add_generic_data(self, worksheet: Worksheet, generic_fields: list[ExtractedField]) -> None:
        worksheet.append(self.GENERIC_DATA_TABLE_HEADER)
        self._add_bold_style_to_row(worksheet, worksheet.max_row)
        for field in generic_fields:
            worksheet.append(
                [
                    field.name,
                    field.page,
                    self.TYPES_REPRESENTATION[field.type],
                    field.data.value,
                    get_confidence_table_representation(field.data.confidence),
                    field.data.validation_result.value,
                ],
            )

    def add_generic_list_fields_data(self, worksheet: Worksheet, generic_list_fields: list[ExtractedField]) -> None:
        for generic_list_field in generic_list_fields:
            worksheet.append(
                self._create_list_common_header(
                    generic_list_field.name,
                    generic_list_field.page,
                    self.TYPES_REPRESENTATION[generic_list_field.base_type],
                ),
            )
            self._add_bold_style_to_odd_cells(worksheet, worksheet.max_row)
            worksheet.append(self.LIST_GENERIC_DATA_TABLE_HEADER)
            self._add_bold_style_to_row(worksheet, worksheet.max_row)
            for index, item in enumerate(generic_list_field.data, 1):
                worksheet.append(
                    [
                        index,
                        item.value,
                        get_confidence_table_representation(item.confidence),
                        item.validation_result.value,
                    ],
                )
            worksheet.append([])
