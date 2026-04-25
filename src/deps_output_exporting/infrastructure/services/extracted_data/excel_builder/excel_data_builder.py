from io import BytesIO

from openpyxl.workbook import Workbook

from deps_output_exporting.infrastructure.proxies import (
    GENERIC_TYPES,
    ExtractedField,
    FieldType,
)

from ..general_info import GeneralInfo
from .general_info_filler import GeneralInfoFiller
from .generic_data_filler import GenericDataFiller
from .kv_data_filler import KeyValueDataFiller
from .table_data_filler import TableDataFiller

__all__ = ["ExcelDataBuilder"]


class ExcelDataBuilder:
    MIN_COLUMN_WIDTH = 6
    MAX_COLUMN_WIDTH = 150

    GENERAL_INFO_SHEET_NAME = "General Information"
    GENERIC_DATA_SHEET_NAME = "Generic Data"
    KV_DATA_SHEET_NAME = "Key Value Pairs"
    TABLE_VIEW_SHEET_NAME = "Table View"
    TABLE_DETAIL_SHEET_NAME = "Table Detail"
    LIST_DATA_SHEET_NAME = "List Data"

    def __init__(self, general_info: GeneralInfo, fields_data: list[ExtractedField]) -> None:
        self._general_info_filler: GeneralInfoFiller = GeneralInfoFiller()
        self._generic_data_filler: GenericDataFiller = GenericDataFiller()
        self._kv_data_filler: KeyValueDataFiller = KeyValueDataFiller()
        self._table_data_filler: TableDataFiller = TableDataFiller()

        self._general_info = general_info

        self._generic_fields: list[ExtractedField] = []
        self._kv_fields: list[ExtractedField] = []
        self._table_fields: list[ExtractedField] = []
        self._generic_list_fields: list[ExtractedField] = []
        self._kv_list_fields: list[ExtractedField] = []
        self._table_list_fields: list[ExtractedField] = []

        self._fill_field_lists(fields_data)

    @classmethod
    def with_info(cls, general_info: GeneralInfo, fields_data: list[ExtractedField]) -> "ExcelDataBuilder":
        return cls(general_info, fields_data)

    def build(self) -> bytes:
        workbook = Workbook()
        active = workbook.active
        workbook.remove(active)

        self._create_general_info_sheet(workbook)
        if self._generic_fields:
            self._create_generic_data_sheet(workbook)
        if self._kv_fields:
            self._create_kv_data_sheet(workbook)
        if self._table_fields:
            self._create_table_view_sheet(workbook)
            self._create_table_detail_sheet(workbook)
        if self._generic_list_fields or self._kv_list_fields or self._table_list_fields:
            self._create_list_data_sheet(workbook)

        self._set_column_width(workbook)

        with BytesIO() as buffer:
            workbook.save(buffer)
            return buffer.getvalue()

    def _fill_field_lists(self, fields_data: list[ExtractedField]) -> None:
        for field in fields_data:
            if field.type in GENERIC_TYPES:
                self._generic_fields.append(field)
            elif field.type == FieldType.DICT:
                self._kv_fields.append(field)
            elif field.type == FieldType.TABLE:
                self._table_fields.append(field)
            elif field.type == FieldType.LIST and field.base_type in GENERIC_TYPES:
                self._generic_list_fields.append(field)
            elif field.type == FieldType.LIST and field.base_type == FieldType.DICT:
                self._kv_list_fields.append(field)
            elif field.type == FieldType.LIST and field.base_type == FieldType.TABLE:
                self._table_list_fields.append(field)

    def _create_general_info_sheet(self, workbook: Workbook) -> None:
        worksheet = workbook.create_sheet(self.GENERAL_INFO_SHEET_NAME)
        self._general_info_filler.add_general_info(worksheet, self._general_info)

    def _create_generic_data_sheet(self, workbook: Workbook) -> None:
        worksheet = workbook.create_sheet(self.GENERIC_DATA_SHEET_NAME)
        self._generic_data_filler.add_generic_data(worksheet, self._generic_fields)

    def _create_kv_data_sheet(self, workbook: Workbook) -> None:
        worksheet = workbook.create_sheet(self.KV_DATA_SHEET_NAME)
        self._kv_data_filler.add_kv_data(worksheet, self._kv_fields)

    def _create_table_view_sheet(self, workbook: Workbook) -> None:
        worksheet = workbook.create_sheet(self.TABLE_VIEW_SHEET_NAME)
        self._table_data_filler.fill_table_view_sheet(worksheet, self._table_fields)

    def _create_table_detail_sheet(self, workbook: Workbook) -> None:
        worksheet = workbook.create_sheet(self.TABLE_DETAIL_SHEET_NAME)
        self._table_data_filler.fill_table_detail_sheet(worksheet, self._table_fields)

    def _create_list_data_sheet(self, workbook: Workbook) -> None:
        worksheet = workbook.create_sheet(self.LIST_DATA_SHEET_NAME)
        self._generic_data_filler.add_generic_list_fields_data(worksheet, self._generic_list_fields)
        self._kv_data_filler.add_kv_list_fields_data(worksheet, self._kv_list_fields)
        self._table_data_filler.add_table_list_fields_data(worksheet, self._table_list_fields)

    def _set_column_width(self, workbook: Workbook) -> None:
        for worksheet in workbook.worksheets:
            for column_cells in worksheet.columns:
                content_length_iterator = (len(str(value)) for cell in column_cells if (value := cell.value))
                try:
                    max_content_length = max(content_length_iterator)
                except ValueError:
                    max_content_length = self.MIN_COLUMN_WIDTH
                length_to_set = min(max_content_length, self.MAX_COLUMN_WIDTH)
                length_to_set = max(length_to_set, self.MIN_COLUMN_WIDTH)
                column_letter = column_cells[0].coordinate[0]
                worksheet.column_dimensions[column_letter].width = length_to_set
