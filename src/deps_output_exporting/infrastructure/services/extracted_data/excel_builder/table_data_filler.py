from openpyxl.styles import Font, colors
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet

from deps_output_exporting.infrastructure.proxies import (
    Cell,
    ExtractedField,
    TableData,
    ValidationResult,
)

from .base_data_filler import BaseDataFiller
from .confidence_calculator import get_confidence_table_representation

__all__ = ["TableDataFiller"]


class TableDataFiller(BaseDataFiller):
    LIST_TABLE_DATA_TABLE_HEADER = ["Index", "Table view", "Table detail"]
    TABLE_DETAIL_TABLE_HEADER = ["Value", "Confidence", "Validation", "Column index", "Row index"]

    def fill_table_view_sheet(self, worksheet: Worksheet, table_fields: list[ExtractedField]) -> None:
        for field in table_fields:
            worksheet.append(self._create_table_common_header(field.name, field.page))
            self._add_bold_style_to_odd_cells(worksheet, worksheet.max_row)
            self._add_table_view(worksheet, field.data, worksheet.max_row + 1)
            worksheet.append([])

    def fill_table_detail_sheet(self, worksheet: Worksheet, table_fields: list[ExtractedField]) -> None:
        for field in table_fields:
            worksheet.append(self._create_table_common_header(field.name, field.page))
            self._add_bold_style_to_odd_cells(worksheet, worksheet.max_row)
            if self._has_required_field_validation(field.data.cells) or self._has_failed_validation(field.data.cells):
                self._add_failed_validation_row(worksheet)
            self._add_table_detail(worksheet, field.data.cells)
            worksheet.append([])

    def add_table_list_fields_data(self, worksheet: Worksheet, table_list_fields: list[ExtractedField]) -> None:
        for table_list_field in table_list_fields:
            links_coordinates: dict[int, tuple[str, str]] = {}
            worksheet.append(
                self._create_list_common_header(
                    table_list_field.name,
                    table_list_field.page,
                    self.TYPES_REPRESENTATION[table_list_field.base_type],
                ),
            )
            self._add_bold_style_to_odd_cells(worksheet, worksheet.max_row)
            all_cells = [cell for table_data in table_list_field.data for cell in table_data.cells]
            if self._has_required_field_validation(all_cells) or self._has_failed_validation(all_cells):
                self._add_failed_validation_row(worksheet)
            worksheet.append(self.LIST_TABLE_DATA_TABLE_HEADER)
            self._add_bold_style_to_row(worksheet, worksheet.max_row)

            info_start_row = worksheet.max_row + 1
            for index, _ in enumerate(table_list_field.data):
                links_coordinates[index] = self._add_item_info(worksheet, info_start_row, index)

            for index, table_data in enumerate(table_list_field.data):
                table_view_range = self._add_item_table_view(worksheet, table_data)
                table_view_name = self._create_excel_table_name(table_list_field.name, index, "view")
                excel_table = self._add_excel_table(worksheet, table_view_name, table_view_range)
                self._add_link_to_table(worksheet, excel_table.ref, links_coordinates[index][0])

            for index, table_data in enumerate(table_list_field.data):
                table_detail_range = self._add_item_table_detail(worksheet, table_data)
                table_detail_name = self._create_excel_table_name(table_list_field.name, index, "detail")
                excel_table = self._add_excel_table(worksheet, table_detail_name, table_detail_range)
                self._add_link_to_table(worksheet, excel_table.ref, links_coordinates[index][1])

        self._unmerge_cells_for_excel_table_format(worksheet)

    def _add_item_info(self, worksheet: Worksheet, info_start_row: int, index: int) -> tuple[str, str]:
        worksheet.cell(row=info_start_row + index, column=1, value=index + 1)
        table_view_cell = worksheet.cell(row=info_start_row + index, column=2)
        table_detail_cell = worksheet.cell(row=info_start_row + index, column=3)
        return (table_view_cell.coordinate, table_detail_cell.coordinate)

    def _add_item_table_view(self, worksheet: Worksheet, table_data: TableData) -> str:
        worksheet.append([])
        start_row = worksheet.max_row + 2
        self._add_table_view(worksheet, table_data, start_row)
        end_row = worksheet.max_row
        return str(
            CellRange(min_col=1, min_row=start_row, max_col=table_data.column_count, max_row=end_row),
        )

    def _add_item_table_detail(self, worksheet: Worksheet, table_data: TableData) -> str:
        worksheet.append([])
        start_row = worksheet.max_row + 2
        self._add_table_detail(worksheet, table_data.cells)
        end_row = worksheet.max_row
        return str(CellRange(min_col=1, min_row=start_row, max_col=5, max_row=end_row))

    def _add_table_view(self, worksheet: Worksheet, table_data: TableData, start_row: int) -> None:
        if table_data.header_row:
            worksheet.append(table_data.header_row)
            self._add_bold_style_to_row(worksheet, worksheet.max_row)
            start_row = worksheet.max_row + 1

        for cell in table_data.cells:
            ws_cell = worksheet.cell(
                row=cell.coordinates.row_index + start_row,
                column=cell.coordinates.column_index + 1,
                value=cell.value,
            )
            if (rowspan := cell.coordinates.rowspan) > 1 and (colspan := cell.coordinates.colspan) > 1:
                worksheet.merge_cells(
                    start_row=ws_cell.row,
                    start_column=ws_cell.column,
                    end_row=ws_cell.row + (rowspan - 1),
                    end_column=ws_cell.column + (colspan - 1),
                )
            elif (rowspan := cell.coordinates.rowspan) > 1:
                worksheet.merge_cells(
                    start_row=ws_cell.row,
                    start_column=ws_cell.column,
                    end_row=ws_cell.row + (rowspan - 1),
                    end_column=ws_cell.column,
                )
            elif (colspan := cell.coordinates.colspan) > 1:
                worksheet.merge_cells(
                    start_row=ws_cell.row,
                    start_column=ws_cell.column,
                    end_row=ws_cell.row,
                    end_column=ws_cell.column + (colspan - 1),
                )

    def _add_table_detail(self, worksheet: Worksheet, table_cells: list[Cell]) -> None:
        worksheet.append(self.TABLE_DETAIL_TABLE_HEADER)
        self._add_bold_style_to_row(worksheet, worksheet.max_row)

        for cell in table_cells:
            worksheet.append(
                [
                    cell.value,
                    get_confidence_table_representation(cell.confidence),
                    cell.validation_result.value,
                    cell.coordinates.column_index,
                    cell.coordinates.row_index,
                ],
            )

    @staticmethod
    def _create_table_common_header(name: str, page: int) -> list:
        return ["Field name", name, "Page", page]

    def _create_excel_table_name(self, field_name: str, index: int, postfix: str) -> str:
        sanitized_field_name = field_name.replace(" ", "_")
        return f"{sanitized_field_name}_{index}_{postfix}"

    def _add_excel_table(self, worksheet: Worksheet, name: str, table_range: str) -> Table:
        table = Table(displayName=name, ref=table_range, headerRowCount=False)
        worksheet.add_table(table)
        return table

    def _add_link_to_table(self, worksheet: Worksheet, table_ref: str, cell_coords: str) -> None:
        worksheet[cell_coords].hyperlink = Hyperlink(ref=cell_coords, target=f"#'{worksheet.title}'!{table_ref}")
        worksheet[cell_coords] = "Go to the table"
        worksheet[cell_coords].font = Font(color=colors.BLUE, underline="single")

    def _add_failed_validation_row(self, worksheet: Worksheet) -> None:
        worksheet.append(["Validation", ValidationResult.FAILED.value])
        self._add_bold_style_to_odd_cells(worksheet, worksheet.max_row)

    def _has_required_field_validation(self, cells: list[Cell]) -> bool:
        not_applied_cells = [cell for cell in cells if cell.validation_result.value == ValidationResult.NOT_APPLIED]

        for cell in not_applied_cells:
            cell.validation_result = ValidationResult.PASSED

        return bool(not_applied_cells)

    def _has_failed_validation(self, cells: list[Cell]) -> bool:
        return any(cell.validation_result.value == ValidationResult.FAILED for cell in cells)

    def _unmerge_cells_for_excel_table_format(self, worksheet: Worksheet) -> None:
        merged_cells = list(worksheet.merged_cells.ranges)
        for merge_range in merged_cells:
            worksheet.unmerge_cells(str(merge_range))
